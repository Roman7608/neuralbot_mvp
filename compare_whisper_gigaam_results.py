#!/usr/bin/env python3
"""
Читает Analytic/SALE/whisper_vs_gigaam.xlsx (лист «Сравнение»),
сравнивает колонки Whisper, GigaAM, Гибрид, Гибрид умный по приветствию после гудков
и предыдущему выводу, записывает отчёт в лист «Вывод».

Запуск после завершения run_sale_whisper_vs_gigaam.py:
  cd /home/romandemo/Vikingi && .venv/bin/python compare_whisper_gigaam_results.py
"""

import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

PROJECT_DIR = Path(__file__).resolve().parent
SALE_DIR = PROJECT_DIR / "Analytic" / "SALE"
XLSX_PATH = SALE_DIR / "whisper_vs_gigaam.xlsx"

# Файлы с переводом (приветствие менеджера после гудков)
FILES_WITH_TRANSFER = (4, 6, 8, 9, 10)

# Паттерн приветствия после перевода: «здравствуйте … менеджер», «менеджер … здравствуйте», «это [имя], менеджер» (файл 6)
_GREETING_PATTERN_A = re.compile(
    r"(меня\s+зовут|здравствуйте|добрый\s+день).*?(менеджер|отдел\s+продаж|викинги|заставн|дилерский|чери)",
    re.IGNORECASE | re.DOTALL,
)
_GREETING_PATTERN_B = re.compile(
    r"(менеджер|отдел\s+продаж|викинги|заставн|дилерский|чери).*?(меня\s+зовут|здравствуйте|добрый\s+день)",
    re.IGNORECASE | re.DOTALL,
)
_GREETING_PATTERN_C = re.compile(
    r"это\s+[а-яё]+[,.\s].*?(менеджер|отдел\s+продаж)",
    re.IGNORECASE | re.DOTALL,
)


def has_greeting_after_transfer(text: str) -> bool:
    if not text or len(text) < 20:
        return False
    # Ищем участок после «переключаю» / «соединяю» (перевод с администратора)
    for sep in ["переключаю", "соединяю", "переведу"]:
        idx = text.lower().find(sep)
        if idx == -1:
            continue
        block = text[idx : idx + 600]
        if (
            _GREETING_PATTERN_A.search(block)
            or _GREETING_PATTERN_B.search(block)
            or _GREETING_PATTERN_C.search(block)
        ):
            return True
    # Исходящий звонок: приветствие в начале (файл 6 — «Наталья, это Андрей, менеджер отдела продаж»)
    if (
        _GREETING_PATTERN_A.search(text[:800])
        or _GREETING_PATTERN_B.search(text[:800])
        or _GREETING_PATTERN_C.search(text[:800])
    ):
        return True
    return False


def snippet_after_transfer(text: str, max_len: int = 220) -> str:
    for sep in ["переключаю", "соединяю", "переведу"]:
        idx = text.lower().find(sep)
        if idx == -1:
            continue
        block = text[idx : idx + max_len]
        return block.strip() or ""
    return ""


def main() -> int:
    if not XLSX_PATH.exists():
        print(f"Файл не найден: {XLSX_PATH}")
        return 1

    df = pd.read_excel(XLSX_PATH, sheet_name="Сравнение")
    cols = list(df.columns)
    has_hybrid = "Гибрид (0=W, 1+=G)" in cols
    has_hybrid_smart = "Гибрид умный" in cols

    # Оценка по файлам с переводом
    lines = [
        "=== Сравнение результатов (Whisper / GigaAM / Гибрид) ===",
        "",
        "Проверка: есть ли приветствие менеджера после «переключаю»/«соединяю» в файлах 4, 6, 8, 9, 10.",
        "",
    ]

    for _, row in df.iterrows():
        fname = row.get("Файл", "")
        try:
            num = int(Path(fname).stem)
        except ValueError:
            continue
        if num not in FILES_WITH_TRANSFER:
            continue

        w = str(row.get("Транскрипт Whisper", "") or "")
        g = str(row.get("Транскрипт GigaAM", "") or "")
        h = str(row.get("Гибрид (0=W, 1+=G)", "") or "") if has_hybrid else ""
        hs = str(row.get("Гибрид умный", "") or "") if has_hybrid_smart else ""

        gw = has_greeting_after_transfer(w)
        gg = has_greeting_after_transfer(g)
        gh = has_greeting_after_transfer(h) if h else False
        ghs = has_greeting_after_transfer(hs) if hs else False

        lines.append(f"--- {fname} ---")
        lines.append(f"  Whisper: приветствие после перевода = {'да' if gw else 'нет'}")
        lines.append(f"  GigaAM:  приветствие после перевода = {'да' if gg else 'нет'}")
        if has_hybrid:
            lines.append(f"  Гибрид (0=W,1+=G): приветствие = {'да' if gh else 'нет'}")
        if has_hybrid_smart:
            lines.append(f"  Гибрид умный: приветствие = {'да' if ghs else 'нет'}")
        lines.append(f"  Длины: W={len(w)}, G={len(g)}" + (f", Гибрид={len(h)}, Умный={len(hs)}" if has_hybrid else ""))
        best_text = hs if (has_hybrid_smart and ghs) else (h if (has_hybrid and gh) else (g if gg else w))
        snip = snippet_after_transfer(best_text, 200)
        if snip:
            lines.append(f"  Фрагмент после перевода: {snip[:180]}...")
        lines.append("")

    # Итог vs предыдущий тест (только файлы 4, 6, 8, 9, 10)
    lines.append("--- Итог vs предыдущий параллельный тест ---")
    lines.append("Ранее: GigaAM давал приветствие в 8 и 10, Whisper — нет; в 4, 6, 9 оба неудовлетворительны.")
    def file_num(r):
        try:
            return int(Path(str(r.get("Файл", ""))).stem)
        except ValueError:
            return 0
    relevant = [r for _, r in df.iterrows() if file_num(r) in FILES_WITH_TRANSFER]
    n_w = sum(1 for r in relevant if has_greeting_after_transfer(str(r.get("Транскрипт Whisper", "") or "")))
    n_g = sum(1 for r in relevant if has_greeting_after_transfer(str(r.get("Транскрипт GigaAM", "") or "")))
    n_h = sum(1 for r in relevant if has_greeting_after_transfer(str(r.get("Гибрид (0=W, 1+=G)", "") or ""))) if has_hybrid else 0
    n_hs = sum(1 for r in relevant if has_greeting_after_transfer(str(r.get("Гибрид умный", "") or ""))) if has_hybrid_smart else 0
    lines.append(f"Сейчас (по файлам 4,6,8,9,10): приветствие есть — Whisper: {n_w}/5, GigaAM: {n_g}/5" + (f", Гибрид: {n_h}/5, Гибрид умный: {n_hs}/5" if has_hybrid else "") + ".")
    if has_hybrid_smart and n_hs >= n_g and n_hs >= n_w:
        lines.append("Рекомендация: использовать «Гибрид умный» в пайплайне ОП.")
    elif has_hybrid and n_h >= n_g:
        lines.append("Рекомендация: использовать «Гибрид (0=W, 1+=G)» в пайплайне ОП.")
    else:
        lines.append("Рекомендация: оставить текущий выбор (Whisper или GigaAM) или донастроить умный выбор.")
    lines.append("")

    report = "\n".join(lines)
    print(report)

    # Записать в лист «Вывод»
    wb = load_workbook(XLSX_PATH)
    if "Вывод" in wb.sheetnames:
        ws = wb["Вывод"]
    else:
        ws = wb.create_sheet("Вывод")
    for row in ws.iter_rows():
        for c in row:
            c.value = None
    for i, line in enumerate(report.splitlines(), 1):
        ws.cell(row=i, column=1, value=line[:32000])
    wb.save(XLSX_PATH)
    print(f"Отчёт записан в лист «Вывод» в {XLSX_PATH}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
